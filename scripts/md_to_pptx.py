#!/usr/bin/env python3
"""Convert a Markdown outline into a PPTX deck with optional speaker notes."""
import argparse
import pathlib
import re
from typing import List, Dict

from pptx import Presentation
from pptx.util import Pt

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^[\-*+]\s+(.*)$")


def parse_markdown_lines(lines: List[str]) -> List[Dict[str, List[str]]]:
    slides: List[Dict[str, List[str]]] = []
    current: Dict[str, List[str]] | None = None

    def ensure_slide(title: str | None = None) -> Dict[str, List[str]]:
        nonlocal current
        if current is None:
            current = {"title": title or f"Slide {len(slides) + 1}", "bullets": [], "notes": []}
        elif title:
            current["title"] = title
        return current

    def finalize_slide():
        nonlocal current
        if current is None:
            return
        if not current["bullets"]:
            current["bullets"].append("(add talking points)")
        slides.append(current)
        current = None

    for raw_line in lines:
        line = raw_line.rstrip("\n")
        stripped = line.strip()
        if not stripped:
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title_text = heading_match.group(2).strip() or f"Slide {len(slides) + 1}"
            if level <= 2:
                finalize_slide()
                ensure_slide(title_text)
                continue
            # lower-level headings become bullets
            ensure_slide()
            ensure_slide()["bullets"].append(title_text)
            continue

        if stripped.startswith(">"):
            note_text = stripped.lstrip("> ")
            ensure_slide()["notes"].append(note_text)
            continue

        bullet_match = BULLET_RE.match(stripped)
        if bullet_match:
            ensure_slide()["bullets"].append(bullet_match.group(1).strip())
            continue

        ensure_slide()["bullets"].append(stripped)

    finalize_slide()
    if not slides:
        slides.append({"title": "Slide 1", "bullets": ["(add content)"] , "notes": []})
    return slides


def slides_to_pptx(slides: List[Dict[str, List[str]]], output_path: pathlib.Path) -> None:
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for slide_data in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data["title"]
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        for idx, bullet in enumerate(slide_data["bullets"]):
            paragraph = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(24)
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.clear()
        if slide_data["notes"]:
            notes_frame.text = "\n".join(slide_data["notes"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=pathlib.Path, help="Path to the source Markdown file")
    parser.add_argument(
        "output",
        nargs="?",
        type=pathlib.Path,
        help="Destination PPTX path (defaults to <markdown_stem>.pptx)",
    )
    args = parser.parse_args()

    markdown_path = args.markdown
    if not markdown_path.exists():
        raise SystemExit(f"Markdown file not found: {markdown_path}")

    output_path = args.output or markdown_path.with_suffix(".pptx")

    slides = parse_markdown_lines(markdown_path.read_text(encoding="utf-8").splitlines())
    slides_to_pptx(slides, output_path)
    print(f"Created {output_path} with {len(slides)} slide(s)")


if __name__ == "__main__":
    main()
